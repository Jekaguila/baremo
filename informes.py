"""Informes descargables: libro de Excel para el docente y PDF individual para el estudiante."""
import io
from datetime import datetime
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics.shapes import Drawing, Line, Rect, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

import baremo as bm
from banco import AREAS, ORDEN_AREAS, PREGUNTAS_POR_ID

# =============================================================================== EXCEL


def _fmt_hoja(ws, nivel_cols=()):
    ws.freeze_panes = "A2"
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for idx, col in enumerate(ws.columns, start=1):
        largo = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(10, largo + 2), 60)
    ws.row_dimensions[1].height = 34
    for idx in nivel_cols:
        for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            for cell in row:
                color = bm.COLORES.get(cell.value)
                if color:
                    cell.fill = PatternFill("solid", fgColor=color.lstrip("#"))
                    cell.font = Font(color="FFFFFF", bold=True)


def excel_resultados(df_res, df_subs, df_int, df_resp, df_est, cfg, items_por_area=None):
    """Devuelve los bytes de un libro de Excel con resultados, subcompetencias, respuestas y baremo."""
    cortes = cfg["cortes"]
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        # ---- Resultados
        res = pd.DataFrame({"Nombre": df_res["nombre"], "Correo": df_res["email"], "Grupo": df_res["grupo"]})
        nivel_cols = []
        for a in ORDEN_AREAS:
            res[f"{AREAS[a]['corto']} (%)"] = df_res[a]
            res[f"{AREAS[a]['corto']} · nivel"] = df_res[f"{a}_nivel"]
            res[f"{AREAS[a]['corto']} · min"] = df_res[f"{a}_min"]
        res["Áreas completadas"] = df_res["completadas"]
        res["Global (%)"] = df_res["global"]
        res["Nivel global"] = df_res["nivel_global"]
        res.to_excel(xw, sheet_name="Resultados", index=False)
        nivel_cols = [i + 1 for i, c in enumerate(res.columns) if "nivel" in c.lower()]
        _fmt_hoja(xw.sheets["Resultados"], nivel_cols)

        # ---- Subcompetencias
        if len(df_subs):
            piv = df_subs.pivot(index="id", columns="subcompetencia", values="pct").round(1)
            orden_cols = [AREAS[a]["subs"][s] for a in ORDEN_AREAS for s in AREAS[a]["subs"]]
            piv = piv.reindex(columns=[c for c in orden_cols if c in piv.columns])
            piv.insert(0, "Nombre", df_res.set_index("id").loc[piv.index, "nombre"])
            piv.reset_index(drop=True).to_excel(xw, sheet_name="Subcompetencias (%)", index=False)
            _fmt_hoja(xw.sheets["Subcompetencias (%)"])

        # ---- Respuestas (detalle)
        if len(df_resp):
            r = df_resp.merge(df_int[["id", "estudiante_id", "area"]], left_on="intento_id", right_on="id")
            r = r.merge(df_est[["id", "nombre", "email"]], left_on="estudiante_id", right_on="id", suffixes=("", "_e"))
            filas = []
            for x in r.itertuples():
                q = PREGUNTAS_POR_ID[x.qid]
                filas.append({
                    "Nombre": x.nombre, "Correo": x.email, "Área": AREAS[x.area]["corto"],
                    "ID pregunta": x.qid, "Subcompetencia": AREAS[x.area]["subs"][q["sub"]],
                    "Respuesta elegida": q["opciones"][x.opcion],
                    "Puntos": bm.puntuar_item(q, x.opcion),
                })
            pd.DataFrame(filas).to_excel(xw, sheet_name="Respuestas", index=False)
            _fmt_hoja(xw.sheets["Respuestas"])

        # ---- Análisis de ítems
        if items_por_area:
            partes = []
            for a, df in items_por_area.items():
                if len(df):
                    d = df.copy()
                    d.insert(0, "Área", AREAS[a]["corto"])
                    partes.append(d)
            if partes:
                pd.concat(partes).to_excel(xw, sheet_name="Análisis de ítems", index=False)
                _fmt_hoja(xw.sheets["Análisis de ítems"])

        # ---- Baremo vigente
        tb = bm.tabla_baremo(cortes)
        tb.to_excel(xw, sheet_name="Baremo vigente", index=False, startrow=0)
        pesos = pd.DataFrame({"Área": [AREAS[a]["nombre"] for a in ORDEN_AREAS],
                              "Ponderación (%)": [cfg["pesos"][a] for a in ORDEN_AREAS]})
        pesos.to_excel(xw, sheet_name="Baremo vigente", index=False, startrow=len(tb) + 3)
        _fmt_hoja(xw.sheets["Baremo vigente"], nivel_cols=[2])
    return buf.getvalue()


# =============================================================================== PDF
_AZUL = colors.HexColor("#1F3A5F")


def _estilos():
    ss = getSampleStyleSheet()
    return {
        "titulo": ParagraphStyle("t", parent=ss["Title"], fontSize=20, textColor=_AZUL, spaceAfter=4),
        "sub": ParagraphStyle("s", parent=ss["Normal"], fontSize=12, textColor=colors.HexColor("#4A5568")),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], fontSize=15, textColor=_AZUL, spaceBefore=12, spaceAfter=6),
        "h3": ParagraphStyle("h3", parent=ss["Heading3"], fontSize=13, textColor=_AZUL, spaceBefore=8, spaceAfter=3),
        "txt": ParagraphStyle("p", parent=ss["Normal"], fontSize=12, leading=16),
        "peq": ParagraphStyle("pq", parent=ss["Normal"], fontSize=10.5, leading=14, textColor=colors.HexColor("#4A5568")),
        "celda": ParagraphStyle("c", parent=ss["Normal"], fontSize=11, leading=14),
        "centro": ParagraphStyle("ce", parent=ss["Normal"], fontSize=11, leading=14, alignment=TA_CENTER),
    }


def _barras(res_areas, cortes, ancho=17 * cm):
    """Gráfico de barras por área, con líneas en los cortes del baremo."""
    alto_fila = 26
    h = alto_fila * len(res_areas) + 36
    d = Drawing(ancho, h)
    x0, wbar = 6.3 * cm, ancho - 6.3 * cm - 1.4 * cm
    for i, (a, pct) in enumerate(res_areas):
        y = h - 36 - i * alto_fila
        d.add(String(0, y + 4, AREAS[a]["corto"], fontName="Helvetica", fontSize=10.5,
                     fillColor=colors.HexColor("#2D3748")))
        d.add(Rect(x0, y, wbar, 14, fillColor=colors.HexColor("#EDF2F7"), strokeColor=None))
        niv = bm.nivel(pct, cortes)
        d.add(Rect(x0, y, wbar * pct / 100.0, 14, fillColor=colors.HexColor(bm.COLORES[niv]), strokeColor=None))
        d.add(String(x0 + wbar + 5, y + 3, f"{pct:.0f} %", fontName="Helvetica-Bold", fontSize=10.5,
                     fillColor=colors.HexColor("#2D3748")))
    for c in cortes:
        x = x0 + wbar * c / 100.0
        d.add(Line(x, 4, x, h - 20, strokeColor=colors.HexColor("#718096"), strokeDashArray=[2, 2], strokeWidth=0.8))
        d.add(String(x - 5, h - 14, str(c), fontName="Helvetica", fontSize=9, fillColor=colors.HexColor("#4A5568")))
    return d


def pdf_individual(estudiante, res_areas, global_pct, cfg, titulo="Informe individual · Baremo de competencias"):
    """PDF con resultados, retroalimentación y baremo para un estudiante.

    res_areas: {area: resultado_area(...)} solo de las áreas completadas.
    global_pct: float o None.
    """
    cortes = cfg["cortes"]
    st = _estilos()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=1.8 * cm, bottomMargin=1.8 * cm, title=titulo,
                            author="Baremo de competencias")
    f = []
    f.append(Paragraph(titulo, st["titulo"]))
    f.append(Paragraph(f"<b>{escape(estudiante['nombre'])}</b> · {escape(estudiante['email'])}", st["sub"]))
    f.append(Paragraph(f"Emitido el {datetime.now():%d/%m/%Y}", st["peq"]))
    f.append(Spacer(1, 10))

    # --- Resultado global
    if global_pct is not None:
        ng = bm.nivel(global_pct, cortes)
        t = Table([[Paragraph(f"<font color='white' size=26><b>{global_pct:.1f}</b></font>"
                              f"<font color='white'> / 100</font>", st["centro"]),
                    Paragraph(f"<b>Nivel global: {ng}</b><br/>{bm.DESCRIPTORES[ng]}", st["celda"])]],
                  colWidths=[5 * cm, 12 * cm])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, 0), colors.HexColor(bm.COLORES[ng])),
            ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor(bm.COLORES[ng])),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 10), ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ]))
        f.append(t)
        f.append(Spacer(1, 4))
        f.append(Paragraph(bm.RETRO_GLOBAL[ng], st["txt"]))
    else:
        f.append(Paragraph("Aún no has completado las cuatro áreas; el resultado global se calculará al terminarlas.", st["txt"]))

    # --- Barras por área
    f.append(Paragraph("Resultado por área", st["h2"]))
    barras = [(a, r["pct"]) for a, r in res_areas.items()]
    if barras:
        f.append(_barras(barras, cortes))
        f.append(Paragraph("Las líneas punteadas marcan los cortes del baremo.", st["peq"]))

    # --- Detalle por área
    for a in ORDEN_AREAS:
        if a not in res_areas:
            continue
        r = res_areas[a]
        n = bm.nivel(r["pct"], cortes)
        bloque = [Paragraph(AREAS[a]["nombre"], st["h3"]),
                  Paragraph(f"<b>{r['pct']:.0f} %</b> ({r['puntos']:.1f} de {r['max']} puntos) · "
                            f"<b>Nivel {n}</b>", st["txt"]),
                  Paragraph(bm.RETRO_AREA[a][n], st["txt"]), Spacer(1, 4)]
        filas = [["Subcompetencia", "Puntaje", "Nivel"]]
        for s, rs in r["subs"].items():
            filas.append([Paragraph(AREAS[a]["subs"][s], st["celda"]), f"{rs['pct']:.0f} %",
                          bm.nivel(rs["pct"], cortes)])
        tab = Table(filas, colWidths=[10.5 * cm, 2.8 * cm, 3.7 * cm])
        est = [("BACKGROUND", (0, 0), (-1, 0), _AZUL), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
               ("FONTSIZE", (0, 0), (-1, -1), 11), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E0")),
               ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (1, 0), (-1, -1), "CENTER")]
        for i, (s, rs) in enumerate(r["subs"].items(), start=1):
            est.append(("BACKGROUND", (2, i), (2, i), colors.HexColor(bm.COLORES[bm.nivel(rs["pct"], cortes)])))
            est.append(("TEXTCOLOR", (2, i), (2, i), colors.white))
        tab.setStyle(TableStyle(est))
        bloque.append(tab)
        _, mejorar = bm.retro_subs(a, r, cortes)
        if mejorar:
            bloque.append(Spacer(1, 4))
            bloque.append(Paragraph("<b>Para fortalecer:</b>", st["txt"]))
            for nom, pct, rec in mejorar:
                bloque.append(Paragraph(f"• <b>{nom}.</b> {rec}", st["txt"]))
        f.append(KeepTogether(bloque))

    # --- Baremo
    f.append(Paragraph("Cómo interpretar los niveles", st["h2"]))
    tb = bm.tabla_baremo(cortes)
    filas = [["Rango (%)", "Nivel", "Descriptor"]]
    for x in tb.itertuples():
        filas.append([x[1], x[2], Paragraph(x[3], st["celda"])])
    tab = Table(filas, colWidths=[2.8 * cm, 3.4 * cm, 10.8 * cm])
    est = [("BACKGROUND", (0, 0), (-1, 0), _AZUL), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
           ("FONTSIZE", (0, 0), (-1, -1), 11), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E0")),
           ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
    for i, nv in enumerate(bm.NIVELES, start=1):
        est.append(("BACKGROUND", (1, i), (1, i), colors.HexColor(bm.COLORES[nv])))
        est.append(("TEXTCOLOR", (1, i), (1, i), colors.white))
    tab.setStyle(TableStyle(est))
    f.append(tab)
    f.append(Spacer(1, 8))
    f.append(Paragraph(bm.AVISO, st["peq"]))
    doc.build(f)
    return buf.getvalue()
